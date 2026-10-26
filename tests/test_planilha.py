import openpyxl

from app.services.planilha import COLUNAS, gerar_modelo


def test_gerar_modelo_tem_cabecalho_na_ordem_esperada():
    arquivo = gerar_modelo()

    wb = openpyxl.load_workbook(arquivo)
    ws = wb.active
    linha_cabecalho = [c.value for c in ws[2]]

    assert linha_cabecalho == COLUNAS


def test_gerar_modelo_tem_linha_de_exemplo_preenchida():
    arquivo = gerar_modelo()

    wb = openpyxl.load_workbook(arquivo)
    ws = wb.active
    linha_exemplo = [c.value for c in ws[3]]

    assert linha_exemplo[COLUNAS.index("PLACA")] == "ABC1234"
    assert all(valor is not None for valor in linha_exemplo)
