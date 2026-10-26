from io import BytesIO

from openpyxl import Workbook, load_workbook

TAMANHO_MAXIMO_BYTES = 5 * 1024 * 1024
LINHAS_MAXIMAS = 500

COLUNAS = [
    "CPF_CNPJ_ASSOCIADO",
    "NUMERO_CONTRATO",
    "PLACA",
    "CHASSI",
    "RENAVAM",
    "MARCA",
    "MODELO",
    "ANO_FABRICACAO",
    "ANO_MODELO",
    "COMBUSTIVEL",
    "TIPO_VEICULO",
    "COR",
    "NOME_PLANO_1",
    "VALOR_PLANO_1",
]

INSTRUCOES = (
    "Apague a linha de exemplo (linha 3) antes de enviar. Preencha uma linha por "
    "veículo, começando na linha 3; todas as colunas são obrigatórias. TIPO_VEICULO "
    "deve ser uma destas categorias: Carros - Motos, Utilitário, PESADO - SEM "
    "RASTREADOR, PESADO - COM RASTREADOR ou Reboque/SemiReboque. NOME_PLANO_1 deve "
    "ser um plano ativo compatível com o TIPO_VEICULO."
)

EXEMPLO = [
    "000.111.222-33",
    "100",
    "ABC1234",
    "9BWZZZ377VT004251",
    "987654321",
    "Fiat",
    "Uno Fire 1.6",
    2010,
    2010,
    "Flex",
    "Carros - Motos",
    "Prata",
    "Plano Prime Drive Completo",
    7.2,
]


def gerar_modelo() -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Modelo"

    ws.append([INSTRUCOES])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLUNAS))

    ws.append(COLUNAS)
    ws.append(EXEMPLO)

    arquivo = BytesIO()
    wb.save(arquivo)
    arquivo.seek(0)
    return arquivo


class PlanilhaInvalida(Exception):
    pass


def ler_planilha(nome_arquivo: str, conteudo: bytes) -> list[dict]:
    if not nome_arquivo.lower().endswith(".xlsx"):
        raise PlanilhaInvalida("Envie um arquivo no formato .xlsx.")

    if len(conteudo) > TAMANHO_MAXIMO_BYTES:
        raise PlanilhaInvalida("O arquivo excede o tamanho máximo de 5 MB.")

    try:
        wb = load_workbook(BytesIO(conteudo), read_only=True, data_only=True)
    except Exception as exc:
        raise PlanilhaInvalida("Não foi possível ler o arquivo. Verifique se é um .xlsx válido.") from exc

    ws = wb.active
    cabecalho = next(ws.iter_rows(min_row=2, max_row=2, values_only=True), None)

    if cabecalho != tuple(COLUNAS):
        raise PlanilhaInvalida("As colunas da planilha não correspondem ao modelo esperado.")

    linhas = [
        {"numero_linha": linha[0].row, "dados": dict(zip(COLUNAS, [celula.value for celula in linha]))}
        for linha in ws.iter_rows(min_row=3)
        if any(celula.value is not None for celula in linha)
    ]

    if len(linhas) > LINHAS_MAXIMAS:
        raise PlanilhaInvalida(f"A planilha tem mais de {LINHAS_MAXIMAS} linhas.")

    return linhas
